import os, json, secrets, threading, time, copy
from datetime import datetime, timezone
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, abort, send_file
import gspread
import qrcode
from io import BytesIO
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from google.oauth2.service_account import Credentials

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY') or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE','true').lower() == 'true', PERMANENT_SESSION_LIFETIME=3600)
LOCK = threading.RLock()
# Reuse the authorized Sheets workbook and short-cache reads to avoid repeated
# requests across ordinary page loads. Writes invalidate the cache immediately.
_BOOK_CACHE_LOCK = threading.RLock()
_DATA_CACHE_LOCK = threading.RLock()
_BOOK_CACHE = None
_SCHEMA_READY = False
_DATA_CACHE = None
_DATA_CACHE_AT = 0.0
try:
    DATA_CACHE_TTL_SECONDS = max(5.0, min(60.0, float(os.environ.get('SHEETS_CACHE_TTL_SECONDS', '20'))))
except ValueError:
    DATA_CACHE_TTL_SECONDS = 20.0
DEFAULT_CATEGORIES = ['Air Scrubbers', 'Dehumidifiers', 'Fans']
LOCATIONS = ['1818', 'Cottontail']
HEADERS = {
 'Inventory': ['Equipment ID','Item Name','Category','Current Status','Current Location','Added At','Added By'],
 'Transactions': ['Transaction ID','Timestamp','Action','Equipment ID','Item Name','Category','Person','Location','Project Number/Name','Requester/Project Manager','Project Location','Unit/Apartment Number','Notes','User'],
 'Categories': ['Category'],
}

def credentials_dict():
    raw = os.environ.get('GOOGLE_SERVICE_ACCOUNT_JSON','').strip()
    if raw:
        return json.loads(raw)
    path = os.environ.get('GOOGLE_APPLICATION_CREDENTIALS','').strip()
    if path:
        with open(path, encoding='utf-8') as f: return json.load(f)
    raise RuntimeError('Google Sheets is not configured. Add GOOGLE_SERVICE_ACCOUNT_JSON to environment variables.')

def get_book():
    global _BOOK_CACHE
    with _BOOK_CACHE_LOCK:
        if _BOOK_CACHE is not None:
            return _BOOK_CACHE
        info = credentials_dict()
        creds = Credentials.from_service_account_info(info, scopes=['https://www.googleapis.com/auth/spreadsheets'])
        client = gspread.authorize(creds)
        sheet_id = os.environ.get('GOOGLE_SHEET_ID', '').strip()
        if not sheet_id:
            raise RuntimeError('GOOGLE_SHEET_ID is missing.')
        _BOOK_CACHE = client.open_by_key(sheet_id)
        return _BOOK_CACHE

def ensure_sheets(book):
    global _SCHEMA_READY
    with _BOOK_CACHE_LOCK:
        if _SCHEMA_READY:
            return
        for title, headers in HEADERS.items():
            try:
                ws = book.worksheet(title)
            except gspread.WorksheetNotFound:
                ws = book.add_worksheet(title=title, rows=1000, cols=max(15, len(headers)))
                ws.append_row(headers)
                values = [headers]
            else:
                values = ws.get_all_values()
                if not values:
                    ws.append_row(headers)
                    values = [headers]
                elif values[0] != headers:
                    raise RuntimeError(f'The {title} sheet has unexpected headers. Expected: {", ".join(headers)}')
            if title == 'Categories' and len(values) == 1:
                for category in DEFAULT_CATEGORIES:
                    ws.append_row([category])
        _SCHEMA_READY = True

def book_ready():
    book = get_book()
    ensure_sheets(book)
    return book

def rows(book, name):
    return book.worksheet(name).get_all_records()

def invalidate_data_cache():
    global _DATA_CACHE, _DATA_CACHE_AT
    with _DATA_CACHE_LOCK:
        _DATA_CACHE = None
        _DATA_CACHE_AT = 0.0

def current_user(): return 'public user'

def get_data():
    global _DATA_CACHE, _DATA_CACHE_AT
    book = book_ready()
    with _DATA_CACHE_LOCK:
        now = time.monotonic()
        if _DATA_CACHE is not None and now - _DATA_CACHE_AT < DATA_CACHE_TTL_SECONDS:
            inventory, transactions, categories = _DATA_CACHE
            return book, copy.deepcopy(inventory), copy.deepcopy(transactions), list(categories)
        inventory = rows(book, 'Inventory')
        transactions = rows(book, 'Transactions')
        categories = [
            str(r.get('Category', '')).strip()
            for r in rows(book, 'Categories')
            if str(r.get('Category', '')).strip()
        ]
        _DATA_CACHE = (inventory, transactions, categories)
        _DATA_CACHE_AT = time.monotonic()
        return book, copy.deepcopy(inventory), copy.deepcopy(transactions), list(categories)

def status_is(item, wanted): return str(item.get('Current Status','')).strip().upper() == wanted

def append_transaction(book, item, action, fields):
    ws = book.worksheet('Transactions')
    now = datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')
    row = [secrets.token_hex(5).upper(), now, action, item.get('Equipment ID',''), item.get('Item Name',''), item.get('Category',''), fields.get('person','').strip(), fields.get('location','').strip(), fields.get('project','').strip(), fields.get('requester','').strip(), fields.get('project_location','').strip(), fields.get('unit','').strip(), fields.get('notes','').strip(), current_user() or 'public employee']
    ws.append_row(row, value_input_option='USER_ENTERED')

@app.context_processor
def inject_globals(): return {'logged_user': current_user(), 'logged_role': session.get('role'), 'locations': LOCATIONS}

@app.route('/manager')
def manager_entry():
    return redirect(url_for('dashboard'))

@app.route('/')
def home(): return render_template('home.html')

@app.route('/manager/dashboard')
def dashboard():
    try:
        _, inventory, transactions, categories = get_data()
        counts = {'total':len(inventory),'in':sum(status_is(x,'IN') for x in inventory),'out':sum(status_is(x,'OUT') for x in inventory),'categories':len(categories)}
        recent = list(reversed(transactions[-10:]))
        return render_template('dashboard.html', counts=counts, recent=recent)
    except Exception as e:
        app.logger.exception('Dashboard load failed'); flash(f'Unable to load Google Sheets data: {e}', 'error'); return render_template('dashboard.html', counts={'total':0,'in':0,'out':0,'categories':0}, recent=[])

@app.route('/movement/<action>', methods=['GET','POST'])
def movement(action):
    if action not in ('in','out'): abort(404)
    try: book, inventory, transactions, categories = get_data()
    except Exception as e:
        app.logger.exception('Data load failed'); flash(f'Unable to connect to Google Sheets. Check deployment configuration. {e}', 'error'); return render_template('movement.html', action=action, items=[], categories=DEFAULT_CATEGORIES)
    # Sign IN must select equipment currently OUT; Sign OUT must select equipment currently IN.
    wanted = 'OUT' if action == 'in' else 'IN'
    available = [i for i in inventory if status_is(i,wanted)]
    if request.method == 'POST':
        ids = list(dict.fromkeys(request.form.getlist('equipment_ids')))
        person = request.form.get('person','').strip(); location=request.form.get('location','').strip()
        if not person or not location or not ids:
            flash('Enter the person, select a location, and select at least one item.', 'error')
        elif location not in LOCATIONS:
            flash('Select a valid location.', 'error')
        else:
            fields={'person':person,'location':location,'project':request.form.get('project',''),'requester':request.form.get('requester',''),'project_location':request.form.get('project_location',''),'unit':request.form.get('unit',''),'notes':request.form.get('notes','')}
            completed=[]; conflicts=[]
            # One Gunicorn worker plus an in-process lock reduces race conditions. For strict multi-user atomicity, use a database.
            with LOCK:
                ws = book.worksheet('Inventory')
                all_values = ws.get_all_values()
                header = all_values[0] if all_values else []
                if not header:
                    raise RuntimeError('The Inventory worksheet has no header row.')
                status_index = header.index('Current Status')
                id_index = header.index('Equipment ID')
                location_col = header.index('Current Location') + 1
                fresh_by_id = {}
                for row_num, values in enumerate(all_values[1:], start=2):
                    record = {key: values[idx] if idx < len(values) else '' for idx, key in enumerate(header)}
                    equipment_id = str(record.get('Equipment ID', '')).strip()
                    if equipment_id:
                        fresh_by_id[equipment_id] = (row_num, record)
                status_col = status_index + 1
                id_col = id_index + 1
                for eid in ids:
                    row_info = fresh_by_id.get(eid)
                    if not row_info:
                        conflicts.append(eid)
                        continue
                    row_num, item = row_info
                    if not status_is(item, wanted):
                        conflicts.append(eid)
                        continue
                    new_status='OUT' if action=='out' else 'IN'
                    ws.update_cell(row_num,status_col,new_status)
                    # Assign the current location only when equipment is signed out; clear it when returned.
                    ws.update_cell(row_num,location_col,fields.get('location','').strip() if action=='out' else '')
                    append_transaction(book,item,'SIGN OUT' if action=='out' else 'SIGN IN',fields)
                    completed.append(eid)
            if completed:
                invalidate_data_cache()
                flash(f'{len(completed)} item(s) successfully signed {"out" if action=="out" else "in"}.', 'success')
            if conflicts:
                invalidate_data_cache()
                flash('Some items were not processed because their status changed. Refresh and try again: '+', '.join(conflicts), 'error')
            return redirect(url_for('movement',action=action))
    return render_template('movement.html', action=action, items=available, categories=categories)

@app.route('/inventory', methods=['GET','POST'])
def inventory_add():
    try: book, inventory, transactions, categories = get_data()
    except Exception as e:
        flash(f'Unable to connect to Google Sheets: {e}', 'error')
        return render_template('inventory.html', categories=DEFAULT_CATEGORIES, inventory=[])
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        category = request.form.get('category', '').strip()
        if not name or category not in categories:
            flash('Enter an item name and select a valid category.', 'error')
        else:
            # Equipment ID is the exact equipment name. Names must be unique
            # because Equipment ID is used as the stable lookup key and QR payload.
            with LOCK:
                fresh = rows(book, 'Inventory')
                duplicate = next((x for x in fresh if
                    str(x.get('Item Name', '')).strip().casefold() == name.casefold()
                    or str(x.get('Equipment ID', '')).strip().casefold() == name.casefold()), None)
                if duplicate:
                    flash('An item with this name or ID already exists. Use a unique equipment name.', 'error')
                else:
                    eid = name
                    book.worksheet('Inventory').append_row(
                        [eid, name, category, 'IN', '', datetime.now().astimezone().isoformat(timespec='seconds'), current_user() or 'public employee'],
                        value_input_option='USER_ENTERED'
                    )
                    invalidate_data_cache()
            if not duplicate:
                flash(f'Equipment added. ID: {eid}', 'success')
                return redirect(url_for('inventory_add', new_id=eid))
    inventory_view = list(reversed(inventory))
    new_id = request.args.get('new_id', '').strip()
    new_item = next((x for x in inventory_view if str(x.get('Equipment ID','')) == new_id), None)
    return render_template('inventory.html', categories=categories, inventory=inventory_view, new_item=new_item)


@app.route('/inventory/qr.png')
def inventory_qr_image():
    """Generate a scannable QR PNG containing the item's Equipment ID."""
    equipment_id = request.args.get('equipment_id', '').strip()
    if not equipment_id or len(equipment_id) > 200:
        abort(404)
    image = qrcode.make(equipment_id)
    output = BytesIO()
    image.save(output, format='PNG')
    output.seek(0)
    safe_name = ''.join(ch if ch.isalnum() or ch in ('-', '_') else '_' for ch in equipment_id)[:80] or 'equipment'
    return send_file(
        output,
        mimetype='image/png',
        as_attachment=request.args.get('download') == '1',
        download_name=f"{safe_name}_QR.png",
        max_age=0,
    )


@app.route('/inventory/qr-label')
def inventory_qr_label():
    """Printable QR label; display fields are passed from the inventory row."""
    equipment_id = request.args.get('equipment_id', '').strip()
    if not equipment_id or len(equipment_id) > 200:
        abort(404)
    item = {
        'Equipment ID': equipment_id,
        'Item Name': request.args.get('name', '').strip() or equipment_id,
        'Category': request.args.get('category', '').strip() or '—',
        'Current Status': request.args.get('status', '').strip() or '—',
        'Current Location': request.args.get('location', '').strip() or '',
    }
    return render_template('qr_label.html', item=item)


@app.route('/inventory/edit', methods=['GET','POST'])
def inventory_edit():
    equipment_id = request.args.get('equipment_id', '').strip()
    try:
        book, inventory, transactions, categories = get_data()
    except Exception as e:
        flash(f'Unable to load inventory: {e}', 'error')
        return redirect(url_for('inventory_add'))
    item = next((x for x in inventory if str(x.get('Equipment ID','')) == equipment_id), None)
    if not item:
        abort(404)
    if request.method == 'POST':
        name = request.form.get('name','').strip()
        category = request.form.get('category','').strip()
        duplicate = next((x for x in inventory if
            str(x.get('Equipment ID','')) != equipment_id and
            (str(x.get('Item Name','')).strip().casefold() == name.casefold()
             or str(x.get('Equipment ID','')).strip().casefold() == name.casefold())), None)
        if not name or category not in categories:
            flash('Enter an item name and select a valid category.', 'error')
        elif duplicate:
            flash('Another item already uses this name or ID. Equipment names must be unique.', 'error')
        else:
            ws = book.worksheet('Inventory')
            values = ws.get_all_values()
            headers = values[0]
            id_col = headers.index('Equipment ID') + 1
            name_col = headers.index('Item Name') + 1
            category_col = headers.index('Category') + 1
            row_num = next((n for n, row in enumerate(values[1:], start=2)
                            if len(row) >= id_col and row[id_col - 1] == equipment_id), None)
            if not row_num:
                flash('Equipment record was not found. Refresh and try again.', 'error')
            else:
                new_id = name
                if new_id != equipment_id:
                    ws.update_cell(row_num, id_col, new_id)
                ws.update_cell(row_num, name_col, name)
                ws.update_cell(row_num, category_col, category)
                if new_id != equipment_id and transactions:
                    tx_ids = [
                        new_id if str(t.get('Equipment ID', '')).strip() == equipment_id
                        else str(t.get('Equipment ID', '')).strip()
                        for t in transactions
                    ]
                    book.worksheet('Transactions').update(
                        f"D2:D{len(tx_ids) + 1}",
                        [[value] for value in tx_ids],
                        value_input_option='RAW'
                    )
                invalidate_data_cache()
                flash(f'Updated equipment. Equipment ID is now: {new_id}.', 'success')
                return redirect(url_for('inventory_add'))
    return render_template('inventory_edit.html', item=item, categories=categories)


@app.route('/transactions')
def transactions_view():
    try: _, inventory, transactions, categories = get_data()
    except Exception as e: flash(f'Unable to load transaction history: {e}', 'error'); transactions=[]
    return render_template('transactions.html', transactions=list(reversed(transactions)))

@app.route('/categories', methods=['GET','POST'])
def categories_view():
    try: book, inventory, transactions, categories = get_data()
    except Exception as e: flash(f'Unable to load categories: {e}', 'error'); categories=DEFAULT_CATEGORIES; book=None
    if request.method == 'POST':
        name=request.form.get('category','').strip()
        if not name: flash('Enter a category name.', 'error')
        elif name.casefold() in [c.casefold() for c in categories]: flash('That category already exists.', 'error')
        elif len(name)>60: flash('Category names must be 60 characters or fewer.', 'error')
        elif book:
            book.worksheet('Categories').append_row([name]); invalidate_data_cache(); flash('Category added.', 'success'); return redirect(url_for('categories_view'))
    return render_template('categories.html', categories=categories)



def _export_rows(kind):
    book, inventory, transactions, categories = get_data()
    if kind == 'inventory':
        headers = HEADERS['Inventory']
        data = [[str(item.get(h, '') or '') for h in headers] for item in inventory]
        title = 'SaLUT Equipment Inventory'
    else:
        headers = HEADERS['Transactions']
        data = [[str(item.get(h, '') or '') for h in headers] for item in transactions]
        title = 'SaLUT Transaction History'
    return headers, data, title

def _excel_export(headers, data, title, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Inventory' if 'Inventory' in title else 'Transactions'
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell = ws.cell(1, 1, title)
    cell.font = Font(name='Aptos Display', size=18, bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='302A27')
    cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    subtitle = ws.cell(2, 1, f"Generated {datetime.now().astimezone().strftime('%d %b %Y, %H:%M')} · SaLUT Equipment Management")
    subtitle.font = Font(italic=True, size=10, color='666666')
    for col, header in enumerate(headers, 1):
        c = ws.cell(4, col, header)
        c.font = Font(bold=True, color='302A27')
        c.fill = PatternFill('solid', fgColor='FFD84A')
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = Border(bottom=Side(style='thin', color='263B7A'))
    for r_idx, row in enumerate(data, 5):
        for c_idx, value in enumerate(row, 1):
            c = ws.cell(r_idx, c_idx, value)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if r_idx % 2 == 0: c.fill = PatternFill('solid', fgColor='F5F6F8')
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(4, 4 + len(data))}"
    for idx, header in enumerate(headers, 1):
        samples = [len(str(header))] + [len(str(row[idx-1] or '')) for row in data[:250]]
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max(samples, default=12) + 2, 13), 32)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    output = BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

def _pdf_export(headers, data, title, filename):
    output = BytesIO()
    page_size = landscape(A4) if len(headers) > 7 else A4
    doc = SimpleDocTemplate(output, pagesize=page_size, rightMargin=10*mm, leftMargin=10*mm, topMargin=12*mm, bottomMargin=12*mm, title=title, author='SaLUT Equipment Management')
    styles = getSampleStyleSheet()
    heading = ParagraphStyle('SaLUTTitle', parent=styles['Title'], textColor=colors.HexColor('#302A27'), fontName='Helvetica-Bold', fontSize=17, alignment=0, spaceAfter=4)
    small = ParagraphStyle('SaLUTSmall', parent=styles['Normal'], textColor=colors.HexColor('#666666'), fontSize=8, spaceAfter=10)
    cell_style = ParagraphStyle('SaLUTCell', parent=styles['Normal'], fontSize=6.5, leading=8, wordWrap='CJK')
    story = [Paragraph(title, heading), Paragraph(f"Generated {datetime.now().astimezone().strftime('%d %b %Y, %H:%M')} · SaLUT Equipment Management", small)]
    table_data = [[Paragraph(str(h).replace('&','&amp;'), cell_style) for h in headers]]
    for row in data:
        table_data.append([Paragraph(str(v or '—').replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'), cell_style) for v in row])
    available_width = page_size[0] - 20*mm
    col_widths = [available_width / len(headers)] * len(headers)
    table = Table(table_data, repeatRows=1, colWidths=col_widths, hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#FFD84A')),('TEXTCOLOR',(0,0),(-1,0),colors.HexColor('#302A27')),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#D9DDE5')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F6F8')])]))
    story.extend([table])
    doc.build(story)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename, mimetype='application/pdf')

@app.route('/export/<kind>/<file_format>')
def export_data(kind, file_format):
    if kind not in ('inventory', 'transactions') or file_format not in ('xlsx', 'pdf'):
        abort(404)
    try:
        headers, data, title = _export_rows(kind)
        # Export only the records matching the search and filters currently applied on screen.
        query = request.args.get('q', '').strip().casefold()
        category = request.args.get('category', '').strip().casefold()
        if query or (kind == 'inventory' and category):
            if kind == 'inventory':
                data = [
                    row for row in data
                    if (not query or query in f"{row[0]} {row[1]}".casefold())
                    and (not category or (len(row) > 2 and str(row[2]).casefold() == category))
                ]
            else:
                # Match the fields displayed in the transaction-history table.
                visible_columns = (1, 2, 3, 4, 5, 6, 7, 8, 9, 13)
                data = [
                    row for row in data
                    if not query or query in ' '.join(
                        str(row[i]) for i in visible_columns if i < len(row)
                    ).casefold()
                ]
            title += ' — Filtered Results'
        filename = f"salut_{kind}_{datetime.now().strftime('%Y%m%d_%H%M')}.{file_format}"
        if file_format == 'xlsx': return _excel_export(headers, data, title, filename)
        return _pdf_export(headers, data, title, filename)
    except Exception as e:
        app.logger.exception('Export failed')
        flash(f'Unable to export {kind}: {e}', 'error')
        return redirect(url_for('inventory_add' if kind == 'inventory' else 'transactions_view'))

@app.route('/health')
def health(): return {'status':'ok'}, 200

@app.errorhandler(403)
def forbidden(e): return render_template('error.html', code=403, message='Manager access is required for this page.'),403
@app.errorhandler(404)
def not_found(e): return render_template('error.html', code=404, message='Page not found.'),404

@app.before_request
def csrf_guard():
    if request.method == 'POST':
        token = session.get('_csrf_token')
        submitted = request.form.get('_csrf_token','')
        if not token or not secrets.compare_digest(token, submitted):
            abort(400)

@app.context_processor
def csrf_context():
    if '_csrf_token' not in session: session['_csrf_token'] = secrets.token_urlsafe(32)
    return {'csrf_token': session['_csrf_token']}

@app.errorhandler(400)
def bad_request(e): return render_template('error.html', code=400, message='Your form expired or could not be verified. Go back, refresh the page, and try again.'),400

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=False)
